import sys
import json
import os
import re

def parse_excel(filepath):
    import openpyxl
    wb = openpyxl.load_workbook(filepath, data_only=True)
    sheet = wb.active
    students = []
    
    for row in sheet.iter_rows(values_only=True):
        if not row: continue
        # Look for a row with a name and gender
        row_strs = [str(cell).strip().lower() if cell is not None else "" for cell in row]
        # Find potential gender column (L/P or Laki-laki/Perempuan)
        gender = None
        for val in row_strs:
            if val in ['l', 'laki-laki', 'laki']: gender = 'Laki-laki'
            elif val in ['p', 'perempuan']: gender = 'Perempuan'
        
        if gender:
            # Name is usually the longest string in the row that doesn't look like a number
            name = ""
            for val in row:
                if val is None: continue
                val_str = str(val).strip()
                if len(val_str) > 3 and not re.match(r'^[\d\W]+$', val_str) and val_str.lower() not in ['l', 'p', 'laki-laki', 'perempuan']:
                    if len(val_str) > len(name):
                        name = val_str
            if name:
                students.append({"name": name, "gender": gender})
                
    return students

def parse_docx(filepath):
    import docx
    doc = docx.Document(filepath)
    students = []
    
    for table in doc.tables:
        for row in table.rows:
            row_data = [cell.text.strip().lower() for cell in row.cells]
            gender = None
            for val in row_data:
                if val in ['l', 'laki-laki', 'laki']: gender = 'Laki-laki'
                elif val in ['p', 'perempuan']: gender = 'Perempuan'
                
            if gender:
                name = ""
                for cell in row.cells:
                    val_str = cell.text.strip()
                    if len(val_str) > 3 and not re.match(r'^[\d\W]+$', val_str) and val_str.lower() not in ['l', 'p', 'laki-laki', 'perempuan']:
                        if len(val_str) > len(name):
                            name = val_str
                if name:
                    students.append({"name": name, "gender": gender})
    return students

def parse_pdf(filepath):
    import fitz # PyMuPDF
    doc = fitz.open(filepath)
    students = []
    
    for page in doc:
        text = page.get_text("text")
        lines = text.split('\n')
        
        # Simple heuristic: look for L or P at the end of lines, or in columns
        # Since PDF text extraction can be messy, let's try to extract lines that look like: "1. Nama Siswa L/P"
        for i, line in enumerate(lines):
            line = line.strip()
            if not line: continue
            
            # Often, tables are extracted line by line. We might have Name, then L/P.
            # Let's check if the next line or this line has L/P
            parts = line.split()
            gender = None
            if len(parts) > 0 and parts[-1].upper() in ['L', 'P']:
                if parts[-1].upper() == 'L': gender = 'Laki-laki'
                else: gender = 'Perempuan'
                name = " ".join(parts[:-1]).strip()
                # Remove leading numbers
                name = re.sub(r'^[\d\.\s]+', '', name)
                if len(name) > 3:
                    students.append({"name": name, "gender": gender})
            else:
                # Check next line for L or P
                if i + 1 < len(lines):
                    next_line = lines[i+1].strip().upper()
                    if next_line in ['L', 'P']:
                        if next_line == 'L': gender = 'Laki-laki'
                        else: gender = 'Perempuan'
                        name = line
                        name = re.sub(r'^[\d\.\s]+', '', name)
                        if len(name) > 3 and not re.match(r'^[\d\W]+$', name):
                            students.append({"name": name, "gender": gender})
    
    # Remove duplicates
    seen = set()
    unique_students = []
    for s in students:
        if s['name'] not in seen:
            seen.add(s['name'])
            unique_students.append(s)
            
    return unique_students

def main():
    files = [
        "data/ABSEN 2A 2026_2027.xlsx",
        "data/ABSEN SISWA 5D 2026 2027.xlsx",
        "data/Absen Kelas 1B.pdf",
        "data/Data siswa 3C 26-27.docx",
        "data/PEROLINGAN KELAS 4A, 4B, dan 4C.pdf",
        "data/Pembagian_Kelas_1_2026_2027.xlsx"
    ]
    
    results = {}
    for f in files:
        if not os.path.exists(f):
            print(f"File not found: {f}", file=sys.stderr)
            continue
        
        ext = f.split('.')[-1].lower()
        try:
            if ext == 'xlsx':
                res = parse_excel(f)
            elif ext == 'docx':
                res = parse_docx(f)
            elif ext == 'pdf':
                res = parse_pdf(f)
            else:
                res = []
            
            results[f] = res
        except Exception as e:
            print(f"Error parsing {f}: {e}", file=sys.stderr)
            
    with open('scripts/parsed_students.json', 'w') as out:
        json.dump(results, out, indent=2)
        
if __name__ == "__main__":
    main()
