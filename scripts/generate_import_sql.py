import openpyxl
import docx
import json
import re

def guess_gender(name):
    n = name.strip().lower()
    last_word = n.split()[-1] if n else ""
    if last_word in ['putri', 'azzahra', 'ayu', 'nisa', 'zahra', 'dewi', 'sari', 'nur', 'sakhi', 'hanun', 'sakina', 'khalisa', 'pustika', 'azalia', 'utami', 'maheswari', 'rahmani', 'wibowo', 'safwana', 'adhisty', 'wimala', 'nistrianto', 'kusuma', 'safira', 'khairunnisa', 'iswahyuni', 'indita', 'darmawan', 'claretta', 'wardhana', 'aya', 'purnomo']:
        return 'Perempuan'
    if last_word in ['putra', 'akbar', 'pratama', 'hidayat', 'ramadhan', 'setiawan', 'prasojo', 'nugroho', 'kamal', 'falah', 'wijaya', 'yusuf', 'wibawa', 'santoso', 'dewantoro', 'wiratama', 'barra', 'siddiq', 'natta', 'annam', 'prawiro', 'pradipta', 'wardana', 'solhidar', 'askari', 'abqari', 'thamrin', 'syah', 'musa', 'ispriono', 'akhyar', 'aliyan', 'alghifari', 'mahindrajaya', 'nurwanto']:
        return 'Laki-laki'
    first_word = n.split()[0] if n else ""
    if first_word in ['aisyah', 'anindita', 'alisha', 'arika', 'azzahra', 'bilqis', 'callista', 'dzakira', 'nafia', 'annisa', 'alika', 'kuntum', 'shevahira', 'tyasayu', 'harumi', 'acintya', 'aifriza', 'ainayya', 'alesha', 'kinarayu', 'gendis', 'adzkiya', 'arsyila', 'elshanum', 'qiandra', 'yurike', 'zizi', 'elmira', 'naila', 'luna', 'nawra', 'assyabia', 'kayyisah', 'shaqueena', 'clarisha', 'chayra', 'adifa', 'adreena', 'aisha', 'almahyra', 'alsyakila', 'aqilla', 'aqila', 'arsya', 'ceisya', 'gladysa', 'khayra', 'myeisha', 'queenzy', 'shakeela', 'zianka', 'adiba', 'bellvania', 'khalifa', 'fiona', 'fifi', 'siti', 'dwi', 'tri', 'ratna', 'clemira', 'enzy', 'kaysha', 'khanza', 'mikhayla', 'nadhira', 'nadine', 'najwa', 'revalena', 'rr', 'r.r']:
        return 'Perempuan'
    return 'Laki-laki'

def clean_name(name):
    if not name: return ""
    name = str(name).strip()
    name = re.sub(r'^\d+[\.\s\-]+', '', name).strip()
    return name

def is_valid_student_name(name):
    if not name or len(name) < 3: return False
    if re.match(r'^[\d\W]+$', name): return False
    bad_kws = ["KOTA", "DAFTAR", "ABSEN", "KELAS", "SEKOLAH", "MADRASAH", "TAHUN", "WALI", "JUMLAH", "MADIUN", "REKAP", "TEMPAT", "TANGGAL", "NAMA", "JENIS", "KETERANGAN", "PEMBAGIAN", "PEMERINTAH", "KEMENTERIAN", "KEPALA", "LAKI", "PEREMPUAN", "TOTAL", "FORMAT", "BULAN", "S.PD", "NIP", "GURU", "NO.", "PAGUYUBAN", "SAMPAH", "PAPAN", "KET", "L/P"]
    if any(kw in name.upper() for kw in bad_kws): return False
    if "…" in name or "..." in name: return False
    return True

# Class targets and their metadata
class_meta = {
    "Kelas 2B": {
        "id": "538cde6c-485e-447c-8fcb-35478f72d7dd",
        "user_id": "e209f787-9a27-4769-9bd8-3761e0617359",
        "file": "data/ABSEN 2B 2627.xlsx"
    },
    "Kelas 2C": {
        "id": "478df0b2-33d8-4e8c-b1dd-78ac5507122c",
        "user_id": "fc5186e1-5551-4a0b-84da-655da7e322b6",
        "file": "data/ABSENSI 2C 2026-2027.docx"
    },
    "Kelas 5A": {
        "id": "609024a9-199b-4615-8a32-65612f5104a3",
        "user_id": "2ab453ab-5ddd-4b19-a5d8-6b2f0f0bf0a6",
        "file": "data/Absensi Siswa Kelas 5A 2627.xlsx"
    },
    "Kelas 5C": {
        "id": "d9a45881-7b02-46d8-a663-c9b8081aa9c8",
        "user_id": "088e7997-3890-4914-bef6-8d084d3f7213",
        "file": "data/DAFTAR PESERTA DIDIK KELAS V-C_T.A 2026-2027.xlsx"
    },
    "Kelas 5D": {
        "id": "a19ae4b7-d562-47ab-ae94-192c80d9427c",
        "user_id": "fc2bfe50-1ad6-4a98-8b04-8e3ca72a163e",
        "file": "data/ABSEN SISWA 5D 2026 2027.xlsx"
    },
    "Kelas 6C": {
        "id": "055d8e8a-0061-401f-8ad8-887b4e319863",
        "user_id": "65f41bf1-4246-4872-9fbf-2f91c229e022",
        "file": "data/ABSENSI 6C 2026-2027.xlsx"
    },
    "Kelas 6D": {
        "id": "a0346e94-b74a-428f-a6f4-0fc11d45241d",
        "user_id": "8b2cb0a4-7a63-4578-bed1-bec4ab8df58d",
        "file": "data/ABSEN 6D BARU 2026.xlsx"
    }
}

clean_predefined = {
    "Kelas 2C": [
        {"name":"ADIFA ARINUR SALAMAH","gender":"Perempuan"},
        {"name":"ADREENA SHEZA KHALISA","gender":"Perempuan"},
        {"name":"AHDA MAUZA","gender":"Laki-laki"},
        {"name":"AISHA DIANDRA PRAMESTI","gender":"Perempuan"},
        {"name":"ALFARIZQI ZAYDAN SISWANDI","gender":"Laki-laki"},
        {"name":"ALMAHYRA SHANUM HIDAYAT","gender":"Perempuan"},
        {"name":"ALSYAKILA GRIZEL NUGRAHA","gender":"Perempuan"},
        {"name":"AQILLA FARIZA MUFIA","gender":"Perempuan"},
        {"name":"AQILA TSURAYYA ACHMAD","gender":"Perempuan"},
        {"name":"ARSYA SAKINA DARIKH","gender":"Perempuan"},
        {"name":"AZKA NARENDRA PUTRA","gender":"Laki-laki"},
        {"name":"BUMI BARAMERU SAHERTIAN","gender":"Laki-laki"},
        {"name":"CEISYA ISMAHENNA SHAQEENA ASHAD","gender":"Perempuan"},
        {"name":"FATIH ADNAN NUGROHO","gender":"Laki-laki"},
        {"name":"FATTAH AZZAMI AHMAD","gender":"Laki-laki"},
        {"name":"GLADYSA RUBY AZKADINA","gender":"Perempuan"},
        {"name":"JOVAN JANITRA PRAMUDI","gender":"Laki-laki"},
        {"name":"KHAYRA AYONA ADZKIYA","gender":"Perempuan"},
        {"name":"M RIDWAN PUTRA HABIBURROHMAN","gender":"Laki-laki"},
        {"name":"MUHAMMAD AZMI AL GHAZALI","gender":"Laki-laki"},
        {"name":"MUHAMMAD HAIKAL ALTHAFAREZKI","gender":"Laki-laki"},
        {"name":"MUHAMMAD KEENAN SETIAWAN","gender":"Laki-laki"},
        {"name":"MYEISHA SHAZIA PUTRI","gender":"Perempuan"},
        {"name":"NAILA AYUDIA CHRISTANTO","gender":"Perempuan"},
        {"name":"QUEENZY ADVINCA LOVABI","gender":"Perempuan"},
        {"name":"RAZKA FAUZAN KAMIL","gender":"Laki-laki"},
        {"name":"SEANDITO RIYANTARA ABDILLAH","gender":"Laki-laki"},
        {"name":"SHAKEELA BETHARI WAHYU ARRAYA","gender":"Perempuan"},
        {"name":"ZAFRAN ARTANABIL","gender":"Laki-laki"},
        {"name":"ZIANKA MAFAZA ELSHANUM PRADIPTA","gender":"Perempuan"}
    ],
    "Kelas 5C": [
        {"name":"AFFAN JULIAN ALFARIZI YUSUF","gender":"Laki-laki"},
        {"name":"AISYAH AQILLA PUTRI","gender":"Perempuan"},
        {"name":"ALIFYA AZZAHRA KHAIRUNNISA","gender":"Perempuan"},
        {"name":"ANDRIANO HAMIZAN PANDYA WARDANA","gender":"Laki-laki"},
        {"name":"ARKHA FAYYADH SISWANDI","gender":"Laki-laki"},
        {"name":"ARSALAN ZIO SOLHIDAR ASKARI","gender":"Laki-laki"},
        {"name":"ARSYAD ZHAFRAN ABQARI","gender":"Laki-laki"},
        {"name":"ATHALLA GIBRAN DWIKA PUTRA","gender":"Laki-laki"},
        {"name":"AYSHA AZZAHRA THAMRIN","gender":"Perempuan"},
        {"name":"BELLVANIA AZZAHRA NISTRIANTO","gender":"Perempuan"},
        {"name":"BHAGAWANTA ALFARIZQI PRADITA","gender":"Laki-laki"},
        {"name":"BONDAN MAHESA ZAFRAN SYAH","gender":"Laki-laki"},
        {"name":"HASNA AKIKO ISWAHYUNI","gender":"Perempuan"},
        {"name":"KAYSHA AZIZA INDITA","gender":"Perempuan"},
        {"name":"KHANZA ARSYILA DARMAWAN","gender":"Perempuan"},
        {"name":"MIKHAYLA ADRIEN CLARETTA","gender":"Perempuan"},
        {"name":"MUHAMMAD MUSA","gender":"Laki-laki"},
        {"name":"MUHAMMAD RAFFASYA ALFARIZQI PRADIPTA","gender":"Laki-laki"},
        {"name":"NABILA SAFIRA KHAIRUNNISA","gender":"Perempuan"},
        {"name":"NADHIRA NUSAIBAH PUTRI ISPRIONO","gender":"Perempuan"},
        {"name":"NADINE RIZKY DEFRIAN AKHYAR","gender":"Perempuan"},
        {"name":"NAJWA AISYAH PUSPITA","gender":"Perempuan"},
        {"name":"NAUFAL ADLAN RAIHAN ALIYAN","gender":"Laki-laki"},
        {"name":"PUTRA SADA ALGHIFARI","gender":"Laki-laki"},
        {"name":"R. REY ARSY MAHINDRAJAYA","gender":"Laki-laki"},
        {"name":"RADITYA SENA NUGRAHA","gender":"Laki-laki"},
        {"name":"RAFARDHAN DHAFIR NURWANTO","gender":"Laki-laki"},
        {"name":"REVALENA AURELLY PUTRI WARDHANA","gender":"Perempuan"},
        {"name":"RR. AIMILIONAMORA ADHITAMA RAYA","gender":"Perempuan"}
    ],
    "Kelas 6C": [
        {"name":"ADELINA DIANDRA WIJAYA","gender":"Perempuan"},
        {"name":"AHNAF HAIDAR FAHMI","gender":"Laki-laki"},
        {"name":"AKMA SHAKILA KHAIRINA","gender":"Perempuan"},
        {"name":"ALI AUNNILLA BAHADJ","gender":"Laki-laki"},
        {"name":"ALVARO CALYA ZIGGY FAUSTAN","gender":"Laki-laki"},
        {"name":"AMALIA PUTRI ANJANI","gender":"Perempuan"},
        {"name":"AMARANGGANA KINNARA PRASETYO","gender":"Perempuan"},
        {"name":"ANOM ANANTA PUTRA","gender":"Laki-laki"},
        {"name":"AZFAR PUTRA YURI","gender":"Laki-laki"},
        {"name":"AZKIYA NAFISA NUGROHO","gender":"Perempuan"},
        {"name":"AZZAHRA RAFANDA ZHAFIRAH","gender":"Perempuan"},
        {"name":"BELAID MUHAMMAD ABD EL RAOUF","gender":"Laki-laki"},
        {"name":"BILQIS ZHARUFA HAWIN AFIQAH","gender":"Perempuan"},
        {"name":"DE ATTAR ALFAREZ FERDIAN","gender":"Laki-laki"},
        {"name":"FATHAN ARSYANENDRA PURNOMO","gender":"Laki-laki"},
        {"name":"HAGIA SOPHIA NAIRA RACHMAN","gender":"Perempuan"},
        {"name":"INA AMINATUS ZAHRO","gender":"Perempuan"},
        {"name":"IRFAN NISMARA SETIAWAN","gender":"Laki-laki"},
        {"name":"KAYLA PRAMUDITA AZALIA","gender":"Perempuan"},
        {"name":"KEANU ABIDZAR RACHMAN","gender":"Laki-laki"},
        {"name":"KEISHA LARASATI RAHARDIYANTO","gender":"Perempuan"},
        {"name":"KHANZA AQILA QOTRUNNADA","gender":"Perempuan"},
        {"name":"MICHELLA AMAIRA RAHMANI KUSUMA","gender":"Perempuan"},
        {"name":"MIKHAEELA ALLESSANDRA SYAHIRA","gender":"Perempuan"},
        {"name":"MUHAMMAD DASTAN AL-KHALEF ANDI MAKMUR","gender":"Laki-laki"},
        {"name":"NUR SYIFA KUSUMA AYU","gender":"Perempuan"},
        {"name":"RAFLI HABIBULLAH GHANI PRASOJO","gender":"Laki-laki"},
        {"name":"REVANO NAUVAL PRATAMA","gender":"Laki-laki"},
        {"name":"RUBY ALIYA DEWINA MARYAM","gender":"Perempuan"},
        {"name":"SAFIA PARAMITA MADINA","gender":"Perempuan"},
        {"name":"SHAIRA HASNA AZKADINA","gender":"Perempuan"},
        {"name":"TALITA HASNA KHUMAIRA","gender":"Perempuan"}
    ]
}

def parse_class_file(c_name, meta):
    if c_name in clean_predefined:
        return clean_predefined[c_name]
    
    filepath = meta["file"]
    ext = filepath.split('.')[-1].lower()
    students = []
    seen = set()

    if ext == 'xlsx':
        wb = openpyxl.load_workbook(filepath, data_only=True)
        sheet = wb.worksheets[0]
        for row in sheet.iter_rows(values_only=True):
            if not row: continue
            row_strs = [str(cell).strip() if cell is not None else "" for cell in row]
            gender = None
            for val in row_strs:
                if val.lower() in ['l', 'laki-laki', 'laki', 'l (laki-laki)']: gender = 'Laki-laki'
                elif val.lower() in ['p', 'perempuan', 'p (perempuan)']: gender = 'Perempuan'
            
            name = ""
            for val in row_strs:
                c_n = clean_name(val)
                if is_valid_student_name(c_n):
                    if len(c_n) > len(name):
                        name = c_n
            if name and name not in seen:
                seen.add(name)
                if not gender: gender = guess_gender(name)
                students.append({"name": name, "gender": gender})
    elif ext == 'docx':
        try:
            doc = docx.Document(filepath)
            for table in doc.tables:
                for row in table.rows:
                    row_strs = [cell.text.strip() for cell in row.cells]
                    gender = None
                    for val in row_strs:
                        if val.lower() in ['l', 'laki-laki', 'laki']: gender = 'Laki-laki'
                        elif val.lower() in ['p', 'perempuan']: gender = 'Perempuan'
                    
                    name = ""
                    for cell in row.cells:
                        c_n = clean_name(cell.text)
                        if is_valid_student_name(c_n):
                            if len(c_n) > len(name):
                                name = c_n
                    if name and name not in seen:
                        seen.add(name)
                        if not gender: gender = guess_gender(name)
                        students.append({"name": name, "gender": gender})
        except Exception:
            pass

    return students

def main():
    all_inserts = []
    for c_name, meta in class_meta.items():
        st_list = parse_class_file(c_name, meta)
        print(f"[{c_name}] -> {len(st_list)} students")
        for s in st_list:
            s_name = s['name'].replace("'", "''")
            s_gender = s['gender']
            class_id = meta['id']
            user_id = meta['user_id']
            sql = f"('{s_name}', '{s_gender}', '{class_id}', '{user_id}')"
            all_inserts.append((c_name, sql))
            
    with open("scripts/import_empty_classes.sql", "w", encoding="utf-8") as f:
        f.write("-- Bulk Insert for empty classes\n")
        by_class = {}
        for c_name, sql in all_inserts:
            by_class.setdefault(c_name, []).append(sql)
            
        for c_name, values_list in by_class.items():
            f.write(f"\n-- {c_name} ({len(values_list)} students)\n")
            f.write("INSERT INTO students (name, gender, class_id, user_id) VALUES\n")
            f.write(",\n".join(values_list) + ";\n")

if __name__ == "__main__":
    main()
